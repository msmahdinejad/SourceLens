from __future__ import annotations

import base64
import json
import re
import shutil
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

ENGINE_ROOT = (
    Path(__file__).parents[1]
    / "src"
    / "paperlens"
    / "renderers"
    / "deck_engine"
)
ENTRYPOINT = ENGINE_ROOT / "entrypoint.mjs"
CONTRACT_VERSION = "paperlens.visual-engine.v1"
LAYOUTS = [
    "cover_cinematic",
    "editorial_split",
    "cards_icon",
    "metrics_dashboard",
    "timeline_story",
    "comparison_dual",
    "chart_focus",
    "matrix_2x2",
    "process_flow",
    "photo_mosaic",
    "quote_photo",
    "big_statement",
    "case_study",
    "table_clean",
    "funnel",
    "dashboard",
    "closing_actions",
]


def _run_engine(request: Path, output: Path, result: Path) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            shutil.which("node") or "node",
            str(ENTRYPOINT),
            "--request",
            str(request),
            "--output",
            str(output),
            "--result",
            str(result),
        ],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
    )


def test_deck_engine_package_declares_the_bridge_contract() -> None:
    package = json.loads((ENGINE_ROOT / "package.json").read_text(encoding="utf-8"))

    assert package["name"] == "@sourcelens/deck-renderer"
    assert package["version"].startswith("5.")
    assert package["paperlens"]["contract"] == CONTRACT_VERSION
    assert package["paperlens"]["layouts"] == 17
    assert package["dependencies"]["pptxgenjs"]


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is not installed")
def test_visual_deck_node_engine_renders_standard_4_by_3_without_overflow(
    tmp_path: Path,
) -> None:
    if not (ENGINE_ROOT / "node_modules" / "pptxgenjs").is_dir():
        pytest.skip("SourceLens deck renderer dependencies are not installed")
    request = tmp_path / "request.json"
    output = tmp_path / "output"
    result = output / "result.json"
    request.write_text(
        json.dumps(
            {
                "contract_version": CONTRACT_VERSION,
                "artifact_type": "slide_deck",
                "artifact_id": "standard-deck",
                "workspace_id": "workspace-test",
                "language": "fa",
                "plan": {
                    "title": "Standard research deck",
                    "slides": [
                        {
                            "layout": "big_statement",
                            "title": "مدل استاندارد",
                            "subtitle": "E = mc²",
                        },
                        {
                            "layout": "closing_actions",
                            "title": "جمع‌بندی",
                            "items": [
                                {"title": "بررسی", "body": "منبع را بررسی کنید."},
                                {"title": "نتیجه", "body": "نتیجه را ثبت کنید."},
                            ],
                        },
                    ],
                },
                "options": {
                    "font_family": "Vazirmatn",
                    "aspect_ratio": "standard",
                    "presenter_name": "پژوهشگر نمونه",
                    "organization": "دانشگاه نمونه",
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    completed = _run_engine(request, output, result)

    assert completed.returncode == 0, completed.stderr
    manifest = json.loads(result.read_text(encoding="utf-8"))
    assert manifest["metadata"]["aspect_ratio"] == "standard"
    assert manifest["metadata"]["canvas_inches"] == {"width": 10, "height": 7.5}
    deck = Presentation(output / "presentation.pptx")
    assert deck.slide_width == Inches(10)
    assert deck.slide_height == Inches(7.5)
    for slide in deck.slides:
        assert all(shape.left + shape.width <= deck.slide_width for shape in slide.shapes)


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is not installed")
def test_deck_engine_entrypoint_renders_editable_chart_notes_and_manifest(
    tmp_path: Path,
) -> None:
    if not (ENGINE_ROOT / "node_modules" / "pptxgenjs").is_dir():
        pytest.skip("SourceLens deck renderer dependencies are not installed")
    request = tmp_path / "request.json"
    output = tmp_path / "output"
    result = output / "result.json"
    request.write_text(
        json.dumps(
            {
                "contract_version": CONTRACT_VERSION,
                "artifact_id": "deck-contract-test",
                "workspace_id": "workspace-test",
                "language": "en",
                "artifact_type": "slide_deck",
                "plan": {
                    "title": "Evidence deck",
                    "subtitle": "Editable output",
                    "theme": {
                        "primary": "173D36",
                        "secondary": "4E7468",
                        "accent": "D6B35D",
                        "accent2": "7CA99A",
                        "background": "F5F1E7",
                        "surface": "FFFCF5",
                        "text": "17211E",
                        "muted": "66716C",
                    },
                    "slides": [
                        {
                            "layout": "cover_cinematic",
                            "mode": "photo",
                            "kicker": "Research brief",
                            "title": "Evidence first",
                            "subtitle": "A grounded narrative",
                            "image_queries": ["research team laboratory"],
                            "speaker_notes": "Introduce the source and scope.",
                        },
                        {
                            "layout": "chart_focus",
                            "mode": "accent",
                            "kicker": "Measured result",
                            "title": "Quality improves",
                            "chart": {
                                "type": "bar",
                                "title": "Answer quality",
                                "labels": ["Baseline", "SourceLens"],
                                "series": [{"name": "Score", "values": [72, 91]}],
                            },
                            "items": [
                                {
                                    "icon": "chart",
                                    "title": "Grounded",
                                    "body": "Every result remains inspectable.",
                                }
                            ],
                            "footnote": "Illustrative test fixture",
                            "citations": ["S1"],
                        },
                        {
                            "layout": "closing_actions",
                            "mode": "dark",
                            "kicker": "Next step",
                            "title": "Start with one question",
                            "items": [
                                {
                                    "icon": "target",
                                    "value": "01",
                                    "title": "Select sources",
                                    "body": "Choose the evidence before generation.",
                                },
                                {
                                    "icon": "chart",
                                    "value": "02",
                                    "title": "Verify values",
                                    "body": "Confirm every numeric claim against its source.",
                                },
                            ],
                        },
                    ],
                },
                "options": {
                    "font_family": "Aptos",
                    "aspect_ratio": "wide",
                    "include_charts": True,
                    "include_speaker_notes": True,
                    "show_image_credits": True,
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    completed = _run_engine(request, output, result)

    assert completed.returncode == 0, completed.stderr
    events = [json.loads(line) for line in completed.stdout.splitlines() if line.strip()]
    assert events[0]["type"] == "progress"
    assert events[0]["stage"] == "normalizing"
    assert events[-1]["stage"] == "completed"
    manifest = json.loads(result.read_text(encoding="utf-8"))
    assert manifest["contract_version"] == CONTRACT_VERSION
    assert manifest["engine"] == {
        "name": "sourcelens-deck-renderer",
        "version": "5.1.0",
    }
    assert manifest["metadata"]["slide_count"] == 3
    assert manifest["metadata"]["layout_distribution"]["chart_focus"] == 1
    assert manifest["metadata"]["editable"] is True
    deck_path = output / manifest["outputs"]["pptx_filename"]
    deck = Presentation(deck_path)
    assert len(deck.slides) == 3
    assert any(shape.has_chart for shape in deck.slides[1].shapes)
    assert "Introduce the source" in deck.slides[0].notes_slide.notes_text_frame.text


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is not installed")
def test_deck_engine_entrypoint_renders_every_layout(tmp_path: Path) -> None:
    if not (ENGINE_ROOT / "node_modules" / "pptxgenjs").is_dir():
        pytest.skip("SourceLens deck renderer dependencies are not installed")
    output = tmp_path / "output"
    result = output / "result.json"
    request = tmp_path / "request.json"
    slides = []
    for index, layout in enumerate(LAYOUTS, start=1):
        slides.append(
            {
                "layout": layout,
                "mode": "photo" if "photo" in layout or layout == "cover_cinematic" else "light",
                "kicker": f"Section {index}",
                "title": f"Layout {layout}",
                "subtitle": "A complete editable layout fixture",
                "body": "One clear message supported by concise evidence.",
                "quote": "Evidence changes the quality of the decision.",
                "quote_source": "SourceLens",
                "callout": "Inspect the evidence before acting.",
                "items": [
                    {
                        "icon": ["target", "chart", "database", "sparkles"][position % 4],
                        "value": f"{position + 1:02d}",
                        "title": f"Grounded dimension {position + 1}",
                        "body": "Grounded supporting detail.",
                    }
                    for position in range(4)
                ],
                "left_title": "Before",
                "right_title": "After",
                "x_axis_label": "Measured impact",
                "y_axis_label": "Implementation effort",
                "left_items": [
                    {"icon": "clock", "title": "Slow", "body": "Manual evidence review."}
                ],
                "right_items": [
                    {"icon": "trend", "title": "Fast", "body": "Inspectable retrieval."}
                ],
                "chart": {
                    "type": "line",
                    "title": "Quality",
                    "labels": ["Before", "After"],
                    "series": [{"name": "Score", "values": [60, 90]}],
                },
                "citations": ["S1"],
                "image_queries": [
                    "research team laboratory",
                    "scientist reviewing evidence",
                    "university data analysis",
                ],
                "speaker_notes": f"Explain layout {layout}.",
            }
        )
    request.write_text(
        json.dumps(
            {
                "contract_version": CONTRACT_VERSION,
                "artifact_type": "slide_deck",
                "artifact_id": "all-layouts",
                "workspace_id": "workspace-test",
                "language": "en",
                "plan": {"title": "All layouts", "slides": slides},
                "options": {
                    "font_family": "Aptos",
                    "aspect_ratio": "wide",
                    "include_charts": True,
                    "include_speaker_notes": True,
                },
            }
        ),
        encoding="utf-8",
    )

    completed = _run_engine(request, output, result)

    assert completed.returncode == 0, completed.stderr
    manifest = json.loads(result.read_text(encoding="utf-8"))
    assert manifest["metadata"]["supported_layouts"] == LAYOUTS
    assert manifest["metadata"]["slide_count"] == 17
    assert set(manifest["metadata"]["layout_distribution"]) == set(LAYOUTS)
    deck_path = output / manifest["outputs"]["pptx_filename"]
    deck = Presentation(deck_path)
    assert len(deck.slides) == 17
    assert sum(any(shape.has_chart for shape in slide.shapes) for slide in deck.slides) == 2
    with ZipFile(deck_path) as archive:
        assert any(name.endswith(".svg") for name in archive.namelist())
        slide_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
    shadows = re.findall(r"<a:outerShdw\b[^>]*>", slide_xml)
    assert shadows
    for element in shadows:
        attributes = dict(re.findall(r'(blurRad|dist|dir)="([^"]+)"', element))
        assert set(attributes) == {"blurRad", "dist", "dir"}
        assert all(value.isdecimal() for value in attributes.values())
        assert int(attributes["blurRad"]) <= 1_000_000
        assert int(attributes["dist"]) <= 1_000_000
        assert int(attributes["dir"]) <= 21_600_000


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is not installed")
def test_deck_engine_never_fabricates_missing_chart_values(tmp_path: Path) -> None:
    if not (ENGINE_ROOT / "node_modules" / "pptxgenjs").is_dir():
        pytest.skip("SourceLens deck renderer dependencies are not installed")
    output = tmp_path / "output"
    result = output / "result.json"
    request = tmp_path / "request.json"
    request.write_text(
        json.dumps(
            {
                "contract_version": CONTRACT_VERSION,
                "artifact_type": "slide_deck",
                "artifact_id": "missing-chart",
                "workspace_id": "workspace-test",
                "language": "fa",
                "plan": {
                    "title": "Evidence only",
                    "slides": [
                        {
                            "layout": "cover_cinematic",
                            "title": "Start",
                            "image_queries": ["research team laboratory"],
                        },
                        {
                            "layout": "chart_focus",
                            "title": "No numeric evidence",
                            "chart": {"labels": ["Only one"], "values": []},
                            "items": [
                                {"title": "Grounded", "body": "No invented values."},
                                {"title": "Traceable", "body": "Use only cited measurements."},
                            ],
                        },
                        {
                            "layout": "closing_actions",
                            "title": "Finish",
                            "items": [
                                {"title": "Check", "body": "Verify the source."},
                                {"title": "Decide", "body": "Use the grounded result."},
                            ],
                        },
                    ],
                },
                "options": {"font_family": "Aptos", "aspect_ratio": "wide"},
            }
        ),
        encoding="utf-8",
    )

    completed = _run_engine(request, output, result)

    assert completed.returncode == 0, completed.stderr
    normalized = json.loads((output / "deck-plan.normalized.json").read_text(encoding="utf-8"))
    repaired = normalized["slides"][1]
    assert repaired["layout"] == "cards_icon"
    assert repaired["chart"] is None
    assert "42" not in json.dumps(repaired, ensure_ascii=False)


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is not installed")
def test_deck_engine_reflows_incomplete_required_media_instead_of_failing(
    tmp_path: Path,
) -> None:
    if not (ENGINE_ROOT / "node_modules" / "pptxgenjs").is_dir():
        pytest.skip("SourceLens deck renderer dependencies are not installed")
    output = tmp_path / "output"
    result = output / "result.json"
    request = tmp_path / "request.json"
    request.write_text(
        json.dumps(
            {
                "contract_version": CONTRACT_VERSION,
                "artifact_type": "slide_deck",
                "artifact_id": "missing-media",
                "workspace_id": "workspace-test",
                "language": "fa",
                "plan": {
                    "title": "Media fallback",
                    "slides": [
                        {
                            "layout": "cover_cinematic",
                            "title": "Grounded cover",
                            "body": "Substantive content remains available without the photo.",
                            "media_required": True,
                        },
                        {
                            "layout": "editorial_split",
                            "title": "Grounded content",
                            "body": "The engine must render this slide without fake media.",
                            "media_required": True,
                        },
                    ],
                },
                "options": {"font_family": "Aptos", "aspect_ratio": "wide"},
            }
        ),
        encoding="utf-8",
    )

    completed = _run_engine(request, output, result)

    assert completed.returncode == 0, completed.stderr
    manifest = json.loads(result.read_text(encoding="utf-8"))
    assert manifest["metadata"]["layout_distribution"] == {"big_statement": 2}
    assert any("verified API media was incomplete" in item for item in manifest["warnings"])


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is not installed")
def test_deck_engine_does_not_embed_local_media_paths(tmp_path: Path) -> None:
    if not (ENGINE_ROOT / "node_modules" / "pptxgenjs").is_dir():
        pytest.skip("SourceLens deck renderer dependencies are not installed")
    media = tmp_path / "private-source.png"
    media.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
    )
    request = tmp_path / "request.json"
    output = tmp_path / "output"
    result = output / "result.json"
    request.write_text(
        json.dumps(
            {
                "contract_version": CONTRACT_VERSION,
                "artifact_type": "slide_deck",
                "artifact_id": "path-privacy",
                "workspace_id": "workspace-test",
                "language": "en",
                "plan": {
                    "title": "Path privacy",
                    "slides": [
                        {
                            "layout": "cover_cinematic",
                            "title": "Embedded media",
                            "media_required": True,
                            "resolved_assets": [{"path": str(media)}],
                        },
                        {"layout": "big_statement", "title": "Summary"},
                    ],
                },
                "options": {},
            }
        ),
        encoding="utf-8",
    )

    completed = _run_engine(request, output, result)

    assert completed.returncode == 0, completed.stderr
    with ZipFile(output / "presentation.pptx") as archive:
        xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        )
    assert str(media.resolve()) not in xml
    assert "private-source.png" not in xml


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is not installed")
def test_deck_engine_uses_cover_identity_and_promotes_equations(
    tmp_path: Path,
) -> None:
    """Regression for user-entered identity being dropped and equations rendered as footnotes."""
    if not (ENGINE_ROOT / "node_modules" / "pptxgenjs").is_dir():
        pytest.skip("SourceLens deck renderer dependencies are not installed")
    output = tmp_path / "output"
    result = output / "result.json"
    request = tmp_path / "request.json"
    equation = r"G_t = \sum_{k=0}^{\infty} \gamma^k r_{t+k+1}"
    request.write_text(
        json.dumps(
            {
                "contract_version": CONTRACT_VERSION,
                "artifact_type": "slide_deck",
                "artifact_id": "identity-equation-regression",
                "workspace_id": "workspace-test",
                "language": "fa",
                "plan": {
                    "title": "Reinforcement learning",
                    "slides": [
                        {
                            "layout": "cover_cinematic",
                            "mode": "photo",
                            "kicker": "Course presentation",
                            "title": "Practical reinforcement learning",
                            "subtitle": "Evidence-backed applications",
                            "image_queries": ["robot learning laboratory"],
                        },
                        {
                            "layout": "big_statement",
                            "mode": "light",
                            "kicker": "02",
                            "title": "Mathematical model",
                            "subtitle": "Value estimation",
                            "body": "The equation is the primary teaching object on this slide.",
                            # Old model output placed equations in footnote; the renderer must
                            # migrate that payload instead of making it unreadably small.
                            "footnote": equation,
                        },
                    ],
                },
                "options": {
                    "font_family": "Aptos",
                    "aspect_ratio": "wide",
                    "presenter_name": "Mina Rahimi",
                    "organization": "University of Isfahan",
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    completed = _run_engine(request, output, result)

    assert completed.returncode == 0, completed.stderr
    normalized = json.loads((output / "deck-plan.normalized.json").read_text(encoding="utf-8"))
    assert normalized["presenterName"] == "Mina Rahimi"
    assert normalized["organization"] == "University of Isfahan"
    assert normalized["slides"][1]["equation"] == equation
    assert normalized["slides"][1]["visualHierarchy"] == "equation"
    assert normalized["slides"][1]["footnote"] == ""

    deck = Presentation(output / "presentation.pptx")
    cover_text = "\n".join(
        shape.text for shape in deck.slides[0].shapes if hasattr(shape, "text")
    )
    assert "Mina Rahimi" in cover_text
    assert "University of Isfahan" in cover_text

    equation_shapes = [
        shape
        for shape in deck.slides[1].shapes
        if hasattr(shape, "text") and "G" in shape.text and "=" in shape.text
    ]
    assert len(equation_shapes) == 1
    equation_runs = [
        run
        for paragraph in equation_shapes[0].text_frame.paragraphs
        for run in paragraph.runs
    ]
    assert equation_runs
    assert min(run.font.size.pt for run in equation_runs if run.font.size) >= 22
    assert all(run.font.bold for run in equation_runs)
    assert all(run.font.color.rgb != RGBColor(255, 255, 255) for run in equation_runs)
    rendered_equation = equation_shapes[0].text
    assert "\\" not in rendered_equation
    assert "{" not in rendered_equation
    assert "∑" in rendered_equation
    assert "∞" in rendered_equation
    assert "γ" in rendered_equation
