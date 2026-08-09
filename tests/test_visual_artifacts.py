import json
import shutil
import threading
import time
from base64 import b64decode
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

from paperlens.fonts import FontRegistry
from paperlens.studio import (
    StudioArtifact,
    StudioArtifactStore,
    StudioService,
    normalize_studio_options,
    studio_creation_catalog,
)
from paperlens.visual_artifacts import (
    GeminiVideoNarrator,
    JamendoMusicResolver,
    MediaAsset,
    PexelsMediaResolver,
    PresentationRenderer,
    VideoOverviewRenderer,
    VisualArtifactService,
)
from paperlens.visual_engine_bridge import VisualEngineCancelled
from paperlens.visual_media import MediaAsset as CanonicalMediaAsset


class _RecordingNarrationRenderer:
    def __init__(self) -> None:
        self.request = None

    async def render(self, request, checkpoint, should_cancel) -> None:
        self.request = request
        for index, _turn in enumerate(request.script, start=1):
            await checkpoint(
                SimpleNamespace(
                    pcm=b"\x10\x27" * 24_000,
                    attempts=[{"attempt": index, "status": "completed"}],
                )
            )


class _MediaPoolStub:
    def __init__(self, responses):
        self.responses = list(responses)
        self.requests = []

    def get_json(self, url, **kwargs):
        self.requests.append((url, kwargs))
        return self.responses.pop(0)


class _PresentationMediaStub:
    available = True

    def photo(
        self,
        _query: str,
        target: Path,
        *,
        orientation: str = "landscape",
        exclude_urls: set[str] | None = None,
    ):
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(
            b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAIAAAABCAQAAABeK7cBAAAADklEQVR42mNk+M8AAAICAQB7CYzhAAAAAElFTkSuQmCC"
            )
        )
        return MediaAsset(target, "photo", "https://example.test/photo")


class _UnavailablePresentationMediaStub:
    available = True

    def photo(
        self,
        _query: str,
        _target: Path,
        *,
        orientation: str = "landscape",
        exclude_urls: set[str] | None = None,
    ):
        raise RuntimeError("Pexels returned no matching licensed asset")


def test_gemini_video_narrator_groups_short_scenes_into_a_natural_live_turn(
    tmp_path: Path,
) -> None:
    artifact = StudioArtifact(
        artifact_id="artifact-video-narration",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="Narrated video",
        status="ready",
        document_ids=["document"],
        language="en",
        options={
            "duration_seconds": 120,
            "narration_model": "gemini-3.1-flash-live-preview",
            "voice": "Kore",
        },
        data={
            "scenes": [
                {"id": "scene-1", "narration": "First bounded scene."},
                {"id": "scene-2", "narration": "Second bounded scene."},
                {"id": "scene-3", "narration": "Third bounded scene."},
            ]
        },
    )
    narrator = GeminiVideoNarrator(
        ["configured-key"],
        turn_timeout_seconds=30,
        minimum_key_gap_seconds=0,
        idle_refresh_seconds=75,
    )
    renderer = _RecordingNarrationRenderer()
    narrator.renderer = renderer

    result = narrator.render(artifact, tmp_path / "narration.wav")

    assert [turn.text for turn in renderer.request.script] == [
        "First bounded scene.\n\nSecond bounded scene.\n\nThird bounded scene."
    ]
    assert result.duration_seconds == pytest.approx(1.0)
    assert len(result.attempts) == 1
    assert [attempt["scene_id"] for attempt in result.attempts] == [
        "narration-block-1",
    ]


def test_video_narrator_supports_piper_only_configuration(tmp_path: Path) -> None:
    class LocalRenderer:
        configured = True
        provider_name = "Piper local TTS"

        async def render(self, request, checkpoint, should_cancel) -> None:
            for turn in request.script:
                await checkpoint(
                    SimpleNamespace(
                        turn_id=turn.turn_id,
                        pcm=b"\x10\x27" * 24_000,
                        attempts=[
                            {
                                "status": "completed",
                                "provider": self.provider_name,
                                "model": "fa_IR-amir-medium.onnx",
                            }
                        ],
                    )
                )

    artifact = StudioArtifact(
        artifact_id="artifact-video-local",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="Local narration",
        status="ready",
        document_ids=["document"],
        language="fa",
        options={"duration_seconds": 60, "voice": "Kore"},
        data={"scenes": [{"id": "scene-1", "narration": "روایت محلی طبیعی."}]},
    )
    narrator = GeminiVideoNarrator(
        [],
        piper_renderer=LocalRenderer(),
        turn_timeout_seconds=30,
        minimum_key_gap_seconds=0,
        idle_refresh_seconds=75,
    )

    result = narrator.render(artifact, tmp_path / "local.wav")

    assert result.duration_seconds == pytest.approx(1.0)
    assert result.attempts[0]["provider"] == "Piper local TTS"


def test_video_narrator_trims_live_turn_silence_and_uses_a_short_natural_gap(
    tmp_path: Path,
) -> None:
    class SilentEdgeRenderer:
        async def render(self, request, checkpoint, should_cancel) -> None:
            self.request = request
            silence = b"\x00\x00" * 12_000
            speech = (b"\x10\x27" * 24_000)
            for index, _turn in enumerate(request.script, start=1):
                await checkpoint(
                    SimpleNamespace(
                        pcm=silence + speech + silence,
                        attempts=[{"attempt": index, "status": "completed"}],
                    )
                )

    artifact = StudioArtifact(
        artifact_id="artifact-video-natural-audio",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="Natural narration",
        status="ready",
        document_ids=["document"],
        language="fa",
        options={"duration_seconds": 60, "voice": "Kore"},
        data={
            "scenes": [
                {"id": "scene-1", "narration": "روایت طبیعی " * 25},
                {"id": "scene-2", "narration": "ادامه خوش‌آهنگ " * 25},
            ]
        },
    )
    narrator = GeminiVideoNarrator(
        ["configured-key"],
        turn_timeout_seconds=30,
        minimum_key_gap_seconds=0,
        idle_refresh_seconds=75,
    )
    renderer = SilentEdgeRenderer()
    narrator.renderer = renderer

    result = narrator.render(artifact, tmp_path / "natural.wav")

    assert 2.1 <= result.duration_seconds <= 2.5
    assert renderer.request.config.pause_milliseconds <= 140
    participant = renderer.request.config.participants[0]
    assert participant.temperature is not None and participant.temperature >= 0.3
    assert "warm" in participant.delivery.casefold()


def test_video_overview_rejects_repeated_narration_before_tts() -> None:
    with pytest.raises(ValueError, match="repeats narration"):
        StudioService._normalize_video_overview(
            {
                "title": "Repeated",
                "scenes": [
                    {
                        "title": "First",
                        "narration": "این همان توضیح ثابت و تکراری برای ویدیو است.",
                    },
                    {
                        "title": "Second",
                        "narration": "این همان توضیح ثابت و تکراری برای ویدیو است!",
                    },
                ],
            }
        )


def test_automatic_video_writer_is_instructed_to_sound_natural() -> None:
    artifact = StudioArtifact(
        artifact_id="artifact-video-natural-script",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="Natural script",
        document_ids=["document"],
        language="fa",
        options=normalize_studio_options("video_overview", {}),
    )

    instructions = StudioService._instructions(artifact).casefold()

    assert "natural spoken" in instructions
    assert "avoid repeated" in instructions


def test_video_font_selection_preserves_the_exact_selected_file(tmp_path: Path) -> None:
    registry = FontRegistry(tmp_path / "fonts")
    bundled = FontRegistry.default_bundled_root()
    registry.add(
        "Vazirmatn-Bold.ttf",
        (bundled / "Vazirmatn" / "Vazirmatn-Bold.ttf").read_bytes(),
    )
    selected = registry.add(
        "Vazirmatn-Regular.ttf",
        (bundled / "Vazirmatn" / "Vazirmatn-Regular.ttf").read_bytes(),
    )

    family, directory, font_file = VideoOverviewRenderer(fonts=registry)._font(
        {"font_family": selected.font_id},
        "متن فارسی",
    )

    assert family == selected.family
    assert directory == selected.path.parent
    assert font_file == selected.path


def _slide_artifact(tmp_path: Path) -> tuple[StudioArtifactStore, StudioArtifact]:
    store = StudioArtifactStore(tmp_path / "workspaces")
    artifact = StudioArtifact(
        artifact_id="artifact-123456abcdef",
        workspace_id="workspace",
        artifact_type="slide_deck",
        title="Grounded deck",
        status="ready",
        document_ids=["document"],
        options=normalize_studio_options(
            "slide_deck",
            {
                "slide_count": 5,
                "image_strategy": "none",
                "aspect_ratio": "wide",
            },
        ),
        data={
            "title": "Grounded deck",
            "slides": [
                {
                    "id": "slide-1",
                    "layout": "cover",
                    "title": "Retrieval, clearly",
                    "subtitle": "A source-grounded overview",
                    "body": "",
                    "bullets": [],
                    "metrics": [],
                    "chart": {},
                    "quote": "",
                    "image_query": "",
                    "speaker_notes": "Introduce the evidence.",
                    "citations": ["S1"],
                },
                {
                    "id": "slide-2",
                    "layout": "key_points",
                    "title": "Key findings",
                    "subtitle": "",
                    "body": "Evidence is retrieved before generation.",
                    "bullets": ["Hybrid search", "Page citations"],
                    "metrics": [],
                    "chart": {},
                    "quote": "",
                    "image_query": "",
                    "speaker_notes": "Close with the main implication.",
                    "citations": ["S1"],
                },
                *[
                    {
                        "id": f"slide-{index}",
                        "layout": "key_points",
                        "title": f"Section {index}",
                        "subtitle": "",
                        "body": "Source-grounded detail.",
                        "bullets": ["Evidence"],
                        "metrics": [],
                        "chart": {},
                        "quote": "",
                        "image_query": "",
                        "speaker_notes": "Explain the source.",
                        "citations": ["S1"],
                    }
                    for index in range(3, 6)
                ],
            ],
        },
    )
    return store, store.save(artifact)


def test_visual_studio_catalog_exposes_real_customization_contracts():
    catalog = {item["artifact_type"]: item for item in studio_creation_catalog()}

    assert catalog["slide_deck"]["fields"]["slide_count"]["maximum"] == 40
    assert catalog["slide_deck"]["defaults"]["slide_count_strategy"] == "auto"
    assert catalog["slide_deck"]["defaults"]["content_depth"] == "auto"
    assert catalog["slide_deck"]["defaults"]["media_density"] == "auto"
    assert catalog["slide_deck"]["defaults"]["design_control"] == "auto"
    assert catalog["slide_deck"]["defaults"]["color_theme"] == "auto"
    assert catalog["slide_deck"]["defaults"]["font_family"] == "auto"
    assert "pexels" in catalog["slide_deck"]["fields"]["image_strategy"]["choices"]
    assert catalog["video_overview"]["fields"]["duration_seconds"]["maximum"] == 1200
    assert catalog["video_overview"]["fields"]["duration_seconds"]["minimum"] == 30
    assert catalog["video_overview"]["defaults"]["duration_seconds"] == 180
    assert catalog["video_overview"]["defaults"]["scene_count_strategy"] == "auto"
    assert "music_usage" not in catalog["video_overview"]["fields"]
    assert catalog["video_overview"]["defaults"]["design_control"] == "auto"
    assert catalog["video_overview"]["defaults"]["pace"] == "auto"
    assert catalog["video_overview"]["defaults"]["transition_style"] == "auto"
    assert catalog["video_overview"]["defaults"]["font_family"] == "auto"
    assert "gemini_voice" in catalog["video_overview"]["fields"]["narration"]["choices"]
    assert (
        "gemini-2.5-flash-native-audio-preview-12-2025"
        in catalog["video_overview"]["fields"]["narration_model"]["choices"]
    )


def test_slide_plan_normalization_rejects_filler_instead_of_padding():
    artifact = StudioArtifact(
        artifact_id="artifact-123456abcdef",
        workspace_id="workspace",
        artifact_type="slide_deck",
        title="Deck",
        document_ids=["document"],
        options=normalize_studio_options("slide_deck", {"slide_count": 5}),
    )
    draft = '{"title":"Deck","slides":[{"title":"Only one"}]}'

    with pytest.raises(ValueError, match="only a title|between 2 and 40"):
        StudioService._decode_output(artifact, draft)


def test_visual_plan_normalizers_preserve_reference_engine_fields():
    deck = StudioService._normalize_slide_deck(
        {
            "title": "Deck",
            "slides": [
                {
                    "title": "Cover",
                    "subtitle": "Grounded overview",
                    "layout": "cover_cinematic",
                    "image_queries": ["research team laboratory"],
                },
                {
                    "title": "Evidence",
                    "layout": "dashboard",
                    "items": [
                        {
                            "title": "Signal",
                            "body": "Measured retrieval quality improved.",
                            "tone": "positive",
                        }
                    ],
                    "citations": ["S1"],
                    "rows": [{"topic": "Recall", "value": "84", "body": "Measured"}],
                    "chart": {
                        "type": "line",
                        "title": "Retrieval quality",
                        "labels": ["A", "B"],
                        "series": [
                            {"name": "Dense", "values": [70, 81]},
                            {"name": "Hybrid", "values": [76, 88]},
                        ],
                    },
                },
                {
                    "title": "Close",
                    "layout": "closing_actions",
                    "items": [
                        {"title": "Verify", "body": "Inspect the cited source before use."},
                        {"title": "Apply", "body": "Use the result in one concrete decision."},
                    ],
                },
            ],
        }
    )
    video = StudioService._normalize_video_overview(
        {
            "title": "Video",
            "scenes": [
                {
                    "title": "Hook",
                    "narration": "A grounded hook.",
                    "icon": "target",
                    "media_policy": ["video", "photo", "card"],
                    "visual": {
                        "layout": "hero_cinematic",
                        "treatment": "cinematic",
                        "typing": {"enabled": True, "cps": 7.5, "variance": 0.2},
                    },
                },
                {
                    "title": "Close",
                    "narration": "A grounded close.",
                    "visual": {"layout": "closing_statement"},
                },
            ],
        }
    )

    assert deck["slides"][1]["rows"][0]["topic"] == "Recall"
    assert deck["slides"][1]["items"][0]["tone"] == "positive"
    assert len(deck["slides"][1]["chart"]["series"]) == 2
    assert deck["slides"][1]["chart"]["title"] == "Retrieval quality"
    assert video["scenes"][0]["icon"] == "target"
    assert video["scenes"][0]["media_policy"] == ["video", "photo", "card"]
    assert video["scenes"][0]["visual"]["typing"]["variance"] == 0.2


def test_presentation_materializer_creates_editable_pptx(tmp_path: Path):
    store, artifact = _slide_artifact(tmp_path)
    service = VisualArtifactService(
        artifacts=store,
        presentations=PresentationRenderer(),
        videos=VideoOverviewRenderer(),
    )

    result = service.materialize(artifact)

    assert result.status == "ready"
    assert result.data["editable"] is True
    assert result.data["direction"]["mode"] == "autonomous"
    assert result.data["slides"][0]["layout"] == "big_statement"
    path = store.asset_path(
        artifact.workspace_id, artifact.artifact_id, result.data["pptx_filename"]
    )
    assert path.is_file()
    deck = Presentation(path)
    assert len(deck.slides) == 5
    assert any(shape.has_text_frame for shape in deck.slides[1].shapes)


def test_presentation_materializer_crops_a_real_media_asset(tmp_path: Path):
    store, artifact = _slide_artifact(tmp_path)
    artifact.options["image_strategy"] = "pexels"
    artifact.data["slides"][0]["image_query"] = "research team laboratory"
    artifact.data["slides"][0]["image_queries"] = ["research team laboratory"]
    artifact = store.save(artifact)
    service = VisualArtifactService(
        artifacts=store,
        presentations=PresentationRenderer(_PresentationMediaStub()),
        videos=VideoOverviewRenderer(),
    )

    result = service.materialize(artifact)

    assert result.status == "ready", result.error
    path = store.asset_path(
        artifact.workspace_id, artifact.artifact_id, result.data["pptx_filename"]
    )
    assert any(shape.shape_type == 13 for shape in Presentation(path).slides[0].shapes)


def test_presentation_materializer_reflows_required_media_when_api_has_no_asset(
    tmp_path: Path,
):
    store, artifact = _slide_artifact(tmp_path)
    artifact.options["image_strategy"] = "pexels"
    artifact.data["slides"][1].update(
        {
            "layout": "editorial_split",
            "image_query": "research team laboratory",
            "image_queries": ["research team laboratory"],
            "media_required": True,
        }
    )
    artifact = store.save(artifact)
    service = VisualArtifactService(
        artifacts=store,
        presentations=PresentationRenderer(_UnavailablePresentationMediaStub()),
        videos=VideoOverviewRenderer(),
    )

    result = service.materialize(artifact)

    assert result.status == "ready", result.error
    assert any(
        "Pexels returned no matching licensed asset" in warning
        for warning in result.data["render_warnings"]
    )
    plan, _options, prepared = PresentationRenderer(
        _UnavailablePresentationMediaStub()
    ).prepare_engine_request(artifact, tmp_path / "engine")
    assert plan["slides"][1]["layout"] == "big_statement"
    assert plan["slides"][1]["media_required"] is False
    assert plan["slides"][1]["image_queries"] == []
    assert any(
        "reflowed from editorial_split" in warning
        for warning in prepared["render_warnings"]
    )


def test_presentation_reflows_when_a_multi_photo_layout_is_only_partially_resolved() -> None:
    slide = {
        "layout": "case_study",
        "media_required": True,
        "image_queries": ["first", "second"],
    }
    warnings: list[str] = []
    one_photo = MediaAsset(Path("one.jpg"), "photo", "https://example.test/one")

    layout = PresentationRenderer()._apply_unresolved_media_fallback(
        slide,
        index=6,
        layout="case_study",
        media_layout="case_study",
        media=[one_photo],
        warnings=warnings,
    )

    assert layout == "big_statement"
    assert slide["media_required"] is False
    assert slide["image_queries"] == []
    assert any("(1/2)" in warning for warning in warnings)


@pytest.mark.parametrize(
    "chart",
    [
        {"type": "bar", "label": "Accuracy", "labels": [], "values": [81.2, 84.7]},
        {"type": "bar", "label": "Accuracy", "labels": [], "values": []},
        {"type": "bar", "label": "Accuracy", "labels": ["A"], "values": []},
        {"type": "bar", "label": "Accuracy", "labels": ["A"], "values": [1, 2]},
        {"type": "bar", "label": "Accuracy", "labels": ["A"], "values": ["high"]},
        {"type": "bar", "label": "Accuracy", "labels": [""], "values": [81.2]},
    ],
)
def test_presentation_materializer_recovers_from_invalid_chart_data(
    tmp_path: Path,
    chart: dict,
):
    store, artifact = _slide_artifact(tmp_path)
    artifact.data["slides"][1].update(
        {
            "layout": "chart",
            "chart": chart,
        }
    )
    artifact = store.save(artifact)
    service = VisualArtifactService(
        artifacts=store,
        presentations=PresentationRenderer(),
        videos=VideoOverviewRenderer(),
    )

    result = service.materialize(artifact)

    assert result.status == "ready", result.error
    assert (
        "Slide 2: chart skipped because categories and numeric values are missing or misaligned"
        in result.data["render_warnings"]
    )
    path = store.asset_path(
        artifact.workspace_id, artifact.artifact_id, result.data["pptx_filename"]
    )
    assert len(Presentation(path).slides) == 5


def test_visual_options_accept_a_user_managed_font_and_keep_auto_as_default():
    automatic = normalize_studio_options("slide_deck", {})
    customized = normalize_studio_options(
        "video_overview",
        {"design_control": "custom", "font_family": "Vazirmatn", "pace": "dynamic"},
    )

    assert automatic["design_control"] == "auto"
    assert automatic["font_family"] == "auto"
    assert customized["font_family"] == "Vazirmatn"
    assert customized["pace"] == "dynamic"


def test_pexels_photo_retries_another_candidate_when_download_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    resolver = PexelsMediaResolver(["test-key"])
    resolver.pool = _MediaPoolStub(
        [
            {
                "photos": [
                    {
                        "width": 3000,
                        "height": 2000,
                        "url": "https://pexels.example/first",
                        "src": {"large": "https://images.example/first.jpg"},
                    },
                    {
                        "width": 2000,
                        "height": 1200,
                        "url": "https://pexels.example/second",
                        "src": {"large": "https://images.example/second.jpg"},
                    },
                ]
            }
        ]
    )
    attempts = []

    def download(url: str, target: Path, *, maximum_bytes: int) -> None:
        attempts.append((url, maximum_bytes))
        if url.endswith("first.jpg"):
            raise RuntimeError("broken asset")
        target.write_bytes(b"image")

    monkeypatch.setattr(PexelsMediaResolver, "_download", staticmethod(download))

    asset = resolver.photo("research", tmp_path / "photo.jpg")

    assert asset.source_url == "https://pexels.example/second"
    assert [item[0] for item in attempts] == [
        "https://images.example/first.jpg",
        "https://images.example/second.jpg",
    ]


def test_pexels_video_selects_smallest_source_that_safely_meets_1080p(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    resolver = PexelsMediaResolver(["test-key"])
    resolver.pool = _MediaPoolStub(
        [
            {
                "videos": [
                    {
                        "duration": 10,
                        "url": "https://www.pexels.com/video/123",
                        "user": {"name": "Real creator"},
                        "video_files": [
                            {
                                "width": 3840,
                                "height": 2160,
                                "link": "https://videos.pexels.com/4k.mp4",
                            },
                            {
                                "width": 1920,
                                "height": 1080,
                                "link": "https://videos.pexels.com/1080p.mp4",
                            },
                        ],
                    }
                ]
            }
        ]
    )
    downloads = []

    def download(url: str, target: Path, *, maximum_bytes: int) -> None:
        downloads.append(url)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"x" * (1_000_000 if "1080p" in url else 4_000_000))

    monkeypatch.setattr(PexelsMediaResolver, "_download", staticmethod(download))

    asset = resolver.video_or_photo(
        "grounded laboratory",
        tmp_path,
        mode="video_first",
        orientation="landscape",
        index=1,
    )

    assert downloads == ["https://videos.pexels.com/1080p.mp4"]
    assert asset.byte_count == 1_000_000


def test_video_media_resolution_propagates_cancellation_into_provider_wait(
    tmp_path: Path,
) -> None:
    checks = 0

    def cancelled() -> bool:
        nonlocal checks
        checks += 1
        return checks >= 3

    class WaitingResolver:
        available = True

        def video_or_photo(self, _query, _target_dir, **kwargs):
            while not kwargs["should_cancel"]():
                time.sleep(0.001)
            raise VisualEngineCancelled("provider wait cancelled")

    artifact = StudioArtifact(
        artifact_id="artifact-cancel12345",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="Cancellation",
        document_ids=["document"],
        options={"media_mode": "video_first"},
    )

    with pytest.raises(VisualEngineCancelled, match="cancelled"):
        VideoOverviewRenderer(media=WaitingResolver())._resolve_media(
            artifact,
            [{"media_required": True, "media_queries": ["licensed research image"]}],
            tmp_path,
            [],
            should_cancel=cancelled,
        )


def test_pexels_cache_replays_bytes_and_real_provenance_without_provider_call(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    cache_dir = tmp_path / "private-cache"
    resolver = PexelsMediaResolver(["test-key"], cache_dir=cache_dir)
    resolver.pool = _MediaPoolStub(
        [
            {
                "photos": [
                    {
                        "width": 1920,
                        "height": 1080,
                        "url": "https://www.pexels.com/photo/123",
                        "photographer": "Real creator",
                        "photographer_url": "https://www.pexels.com/@real",
                        "src": {"large": "https://images.pexels.com/123.jpg"},
                    }
                ]
            }
        ]
    )
    downloads = []

    def download(url: str, target: Path, *, maximum_bytes: int) -> None:
        downloads.append((url, maximum_bytes))
        target.write_bytes(b"licensed-image")

    monkeypatch.setattr(PexelsMediaResolver, "_download", staticmethod(download))
    first = resolver.photo("grounded query", tmp_path / "first.jpg")

    replay = PexelsMediaResolver(["test-key"], cache_dir=cache_dir)
    replay.pool = _MediaPoolStub([])
    second = replay.photo("grounded query", tmp_path / "second.jpg")

    assert not first.cache_hit
    assert second.cache_hit
    assert second.byte_count == len(b"licensed-image")
    assert second.path.read_bytes() == b"licensed-image"
    assert second.credit() == first.credit()
    assert replay.pool.requests == []
    assert len(downloads) == 1


def test_jamendo_music_retries_a_broader_query_when_specific_search_is_empty(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    resolver = JamendoMusicResolver(["test-client"])
    resolver.pool = _MediaPoolStub(
        [
            {"results": []},
            {
                "results": [
                    {
                        "audio": "https://audio.example/no-derivatives.mp3",
                        "shareurl": "https://jamendo.example/no-derivatives",
                        "artist_name": "Restricted artist",
                        "license_ccurl": "https://creativecommons.org/licenses/by-nc-nd/3.0/",
                    },
                    {
                        "audio": "https://audio.example/track.mp3",
                        "shareurl": "https://jamendo.example/track",
                        "artist_name": "Test artist",
                        "license_ccurl": "http://creativecommons.org/licenses/by/4.0/",
                    }
                ]
            },
            {"results": []},
        ]
    )
    monkeypatch.setattr(
        PexelsMediaResolver,
        "_download",
        staticmethod(lambda _url, target, *, maximum_bytes: target.write_bytes(b"audio")),
    )

    asset = resolver.resolve(
        "instrumental very specific research title",
        tmp_path / "music.mp3",
        usage="noncommercial",
    )

    assert asset.path.read_bytes() == b"audio"
    assert asset.license_url == "https://creativecommons.org/licenses/by/4.0/"
    assert [request[1]["params"]["search"] for request in resolver.pool.requests] == [
        "instrumental very specific research title",
        "instrumental ambient",
    ]


def test_presentation_materializer_keeps_a_valid_editable_chart(tmp_path: Path):
    store, artifact = _slide_artifact(tmp_path)
    artifact.data["slides"][1].update(
        {
            "layout": "chart",
            "chart": {
                "type": "column",
                "label": "Accuracy",
                "labels": ["Baseline", "SourceLens"],
                "values": [81.2, 84.7],
            },
        }
    )
    artifact = store.save(artifact)
    service = VisualArtifactService(
        artifacts=store,
        presentations=PresentationRenderer(),
        videos=VideoOverviewRenderer(),
    )

    result = service.materialize(artifact)

    assert result.status == "ready", result.error
    assert not any("chart skipped" in warning for warning in result.data["render_warnings"])
    path = store.asset_path(
        artifact.workspace_id, artifact.artifact_id, result.data["pptx_filename"]
    )
    deck = Presentation(path)
    assert any(shape.has_chart for shape in deck.slides[1].shapes)


def test_presentation_materializer_validates_chart_after_direction_normalization(
    tmp_path: Path,
):
    store, artifact = _slide_artifact(tmp_path)
    artifact.data["slides"][1].update(
        {
            "layout": "chart_focus",
            "citations": ["S1"],
            "chart": {
                "type": "bar",
                "labels": ["Current", "Restored"],
                "series": [
                    {"name": "Damage", "values": ["80", "20"]},
                    {"name": "Benefit", "values": ["0", "60"]},
                ],
            },
        }
    )
    artifact = store.save(artifact)

    result = VisualArtifactService(
        artifacts=store,
        presentations=PresentationRenderer(),
        videos=VideoOverviewRenderer(),
    ).materialize(artifact)

    assert result.status == "ready", result.error
    assert not any("chart skipped" in warning for warning in result.data["render_warnings"])
    path = store.asset_path(
        artifact.workspace_id, artifact.artifact_id, result.data["pptx_filename"]
    )
    assert any(shape.has_chart for shape in Presentation(path).slides[1].shapes)


def test_presentation_renderer_enforces_disabled_chart_notes_and_citations(tmp_path: Path):
    store, artifact = _slide_artifact(tmp_path)
    artifact.options.update(
        {
            "include_charts": False,
            "include_speaker_notes": False,
            "show_citations": False,
        }
    )
    artifact.data["slides"][1].update(
        {
            "layout": "chart",
            "chart": {
                "type": "column",
                "label": "Accuracy",
                "labels": ["Baseline", "SourceLens"],
                "values": [81.2, 84.7],
            },
        }
    )
    artifact = store.save(artifact)
    service = VisualArtifactService(
        artifacts=store,
        presentations=PresentationRenderer(),
        videos=VideoOverviewRenderer(),
    )

    result = service.materialize(artifact)

    assert result.status == "ready", result.error
    path = store.asset_path(
        artifact.workspace_id, artifact.artifact_id, result.data["pptx_filename"]
    )
    slide = Presentation(path).slides[1]
    assert not any(shape.has_chart for shape in slide.shapes)
    assert "S1" not in " ".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert "Close with the main implication." not in slide.notes_slide.notes_text_frame.text


def test_slide_plan_normalization_reflows_empty_chart_payload(tmp_path: Path):
    _, artifact = _slide_artifact(tmp_path)
    artifact.data["slides"][1].update(
        {
            "layout": "chart",
            "chart": {
                "type": "bar",
                "label": "Accuracy",
                "labels": [],
                "values": [],
            },
        }
    )

    data, _markdown, _title = StudioService._decode_output(
        artifact,
        json.dumps(artifact.data),
    )

    assert data["slides"][1]["layout"] == "big_statement"
    assert data["slides"][1]["chart"] == {}


def test_python_visual_renderers_only_prepare_inputs_for_node() -> None:
    presentation = PresentationRenderer()
    video = VideoOverviewRenderer()

    assert not hasattr(presentation, "render")
    assert not hasattr(presentation, "_render_professional_slide")
    assert not hasattr(video, "render")
    assert not hasattr(video, "_render_segment")
    assert not hasattr(video, "_compose_segments")


def test_video_retime_preserves_agent_selected_scene_rhythm() -> None:
    scenes = [
        {"duration_seconds": 5, "narration": "Opening"},
        {"duration_seconds": 15, "narration": "Evidence"},
        {"duration_seconds": 10, "narration": "Conclusion"},
    ]

    VideoOverviewRenderer._retime(scenes, 60)

    assert [scene["duration_seconds"] for scene in scenes] == [10.0, 30.0, 20.0]
    assert sum(scene["duration_seconds"] for scene in scenes) == 60.0


def test_video_engine_request_extends_instead_of_accelerating_long_narration(
    tmp_path: Path,
) -> None:
    class LongerNarrator:
        def render(self, artifact, target):
            target.write_bytes(b"RIFF")
            return SimpleNamespace(path=target, duration_seconds=40.7, attempts=[])

    artifact = StudioArtifact(
        artifact_id="artifact-video-exact-duration",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="Exact duration",
        document_ids=["document"],
        options={
            "duration_seconds": 40,
            "narration": "gemini_voice",
            "media_mode": "motion_graphics",
        },
        data={
            "title": "Exact duration",
            "scenes": [
                {"id": "scene-1", "duration_seconds": 10, "narration": "Opening"},
                {"id": "scene-2", "duration_seconds": 30, "narration": "Evidence"},
            ],
        },
    )

    plan, options, prepared = VideoOverviewRenderer(
        narrator=LongerNarrator()
    ).prepare_engine_request(artifact, tmp_path)

    assert options["duration_seconds"] == pytest.approx(41.2)
    assert sum(scene["duration_seconds"] for scene in plan["scenes"]) == pytest.approx(41.2)
    assert prepared["narration_fit"]["source_duration_seconds"] == pytest.approx(40.7)
    assert prepared["narration_fit"]["target_duration_seconds"] == pytest.approx(41.2)
    assert prepared["narration_fit"]["tempo_ratio"] == 1
    assert prepared["narration_fit"]["speed_changed"] is False


def test_video_materializer_persists_the_narration_adjusted_scene_timeline(
    tmp_path: Path,
) -> None:
    class LongerNarrator:
        def render(self, _artifact, target):
            target.write_bytes(b"natural narration")
            return SimpleNamespace(path=target, duration_seconds=40.7, attempts=[])

    class Bridge:
        def render(self, **kwargs):
            return SimpleNamespace(
                warnings=[],
                artifact_data=lambda: {
                    "duration_seconds": kwargs["options"]["duration_seconds"],
                    "video_filename": "video-overview.mp4",
                },
            )

    store = StudioArtifactStore(tmp_path / "workspaces")
    artifact = store.save(
        StudioArtifact(
            artifact_id="artifact-a1b2c3d4e5f6",
            workspace_id="workspace",
            artifact_type="video_overview",
            title="Natural timing",
            status="ready",
            document_ids=["document"],
            options={
                "duration_seconds": 40,
                "narration": "gemini_voice",
                "media_mode": "motion_graphics",
            },
            data={
                "scenes": [
                    {"id": "scene-1", "duration_seconds": 10, "narration": "Opening"},
                    {"id": "scene-2", "duration_seconds": 30, "narration": "Evidence"},
                ]
            },
        )
    )

    result = VisualArtifactService(
        artifacts=store,
        presentations=PresentationRenderer(),
        videos=VideoOverviewRenderer(narrator=LongerNarrator()),
        bridge=Bridge(),
    ).materialize(artifact)

    assert result.status == "ready"
    assert sum(scene["duration_seconds"] for scene in result.data["scenes"]) == pytest.approx(
        41.2
    )


def test_required_video_media_never_silently_becomes_a_generated_card(tmp_path: Path) -> None:
    artifact = StudioArtifact(
        artifact_id="artifact-video-media-1",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="API media",
        document_ids=["document"],
        options={"media_mode": "auto"},
    )

    with pytest.raises(RuntimeError, match="Verified API media is required"):
        VideoOverviewRenderer()._resolve_media(
            artifact,
            [
                {
                    "media_required": True,
                    "media_queries": ["laboratory researchers using microscope"],
                }
            ],
            tmp_path,
            [],
        )


def test_video_media_resolution_is_bounded_and_concurrent(tmp_path: Path) -> None:
    lock = threading.Lock()
    active = 0
    peak = 0

    class SlowLicensedResolver:
        available = True

        def video_or_photo(self, query, target_dir, **kwargs):
            nonlocal active, peak
            with lock:
                active += 1
                peak = max(peak, active)
            time.sleep(0.08)
            with lock:
                active -= 1
            index = kwargs["index"]
            path = target_dir / f"scene-{index:02d}.jpg"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"licensed")
            return MediaAsset(
                path=path,
                kind="photo",
                source_url=f"https://pexels.example/{index}",
                creator="Fixture creator",
                license_url="https://www.pexels.com/license/",
            )

    artifact = StudioArtifact(
        artifact_id="artifact-media-pool",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="Media pool",
        document_ids=["document"],
        options={"media_mode": "video_first"},
    )
    scenes = [
        {"media_queries": [f"licensed research image {index}"], "duration_seconds": 5}
        for index in range(4)
    ]
    assets = VideoOverviewRenderer(
        media=SlowLicensedResolver(), media_workers=3
    )._resolve_media(artifact, scenes, tmp_path, [])

    assert len(assets) == 4
    assert peak == 3


def test_video_media_resolution_retries_a_duplicate_with_used_urls(tmp_path: Path) -> None:
    calls: list[set[str]] = []

    class DuplicateThenUniqueResolver:
        available = True

        def video_or_photo(self, _query, target_dir, **kwargs):
            excluded = set(kwargs["exclude_urls"])
            calls.append(excluded)
            index = kwargs["index"]
            source_url = (
                f"https://pexels.example/unique-{index}"
                if "https://pexels.example/duplicate" in excluded
                else "https://pexels.example/duplicate"
            )
            path = target_dir / f"scene-{index:02d}.jpg"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"licensed")
            return MediaAsset(
                path=path,
                kind="photo",
                source_url=source_url,
                creator="Fixture creator",
                license_url="https://www.pexels.com/license/",
            )

    artifact = StudioArtifact(
        artifact_id="artifact-media-duplicate",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="Unique media",
        document_ids=["document"],
        options={"media_mode": "video_first"},
    )
    scenes = [
        {"media_required": True, "media_queries": [f"research scene {index}"]}
        for index in range(1, 3)
    ]

    assets = VideoOverviewRenderer(
        media=DuplicateThenUniqueResolver(), media_workers=2
    )._resolve_media(artifact, scenes, tmp_path, [])

    assert [asset.source_url for asset in assets] == [
        "https://pexels.example/duplicate",
        "https://pexels.example/unique-2",
    ]
    assert calls[-1] == {"https://pexels.example/duplicate"}


def test_video_media_resolution_reports_each_scene_live(tmp_path: Path) -> None:
    class UniqueResolver:
        available = True

        def video_or_photo(self, _query, target_dir, **kwargs):
            index = kwargs["index"]
            path = target_dir / f"scene-{index:02d}.jpg"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"licensed")
            return MediaAsset(
                path=path,
                kind="photo",
                source_url=f"https://pexels.example/{index}",
                license_url="https://www.pexels.com/license/",
            )

    artifact = StudioArtifact(
        artifact_id="artifact-media-progress",
        workspace_id="workspace",
        artifact_type="video_overview",
        title="Live progress",
        document_ids=["document"],
        options={"media_mode": "video_first"},
    )
    updates: list[tuple[int, int, int, str]] = []

    VideoOverviewRenderer(media=UniqueResolver(), media_workers=2)._resolve_media(
        artifact,
        [{"media_queries": [f"scene {index}"]} for index in range(1, 4)],
        tmp_path,
        [],
        on_progress=lambda scene, current, total, status: updates.append(
            (scene, current, total, status)
        ),
    )

    assert {update[0] for update in updates} == {1, 2, 3}
    assert updates[-1][1:3] == (3, 3)
    assert all(update[3] == "resolved" for update in updates)


@pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="FFmpeg is not installed")
@pytest.mark.parametrize("include_captions", [True, False])
def test_video_materializer_honors_caption_setting(
    tmp_path: Path,
    include_captions: bool,
):
    store = StudioArtifactStore(tmp_path / "workspaces")
    artifact = store.save(
        StudioArtifact(
            artifact_id="artifact-fedcba654321",
            workspace_id="workspace",
            artifact_type="video_overview",
            title="RAG overview",
            status="ready",
            document_ids=["document"],
            options=normalize_studio_options(
                "video_overview",
                {
                    "duration_seconds": 30,
                    "scene_count_strategy": "fixed",
                    "scene_count": 4,
                    "resolution": "720p",
                    "media_mode": "motion_graphics",
                    "narration": "subtitles",
                    "include_captions": include_captions,
                },
            )
            | {"duration_seconds": 30},
            data={
                "title": "RAG overview",
                "scenes": [
                    {
                        "id": "scene-1",
                        "duration_seconds": 3,
                        "kicker": "01",
                        "title": "The problem",
                        "on_screen_text": "Answers need evidence",
                        "narration": "A reliable answer starts with source evidence.",
                        "media_query": "",
                        "visual_intent": "evidence cards",
                        "transition": "fade",
                        "citations": ["S1"],
                    },
                    {
                        "id": "scene-2",
                        "duration_seconds": 3,
                        "kicker": "02",
                        "title": "The result",
                        "on_screen_text": "Grounded generation",
                        "narration": "Retrieval and citations make the result auditable.",
                        "media_query": "",
                        "visual_intent": "connected nodes",
                        "transition": "fade",
                        "citations": ["S1"],
                    },
                    {
                        "id": "scene-3",
                        "duration_seconds": 3,
                        "kicker": "03",
                        "title": "The workflow",
                        "on_screen_text": "Retrieve, verify, answer",
                        "narration": "The workflow keeps evidence connected to each answer.",
                        "media_query": "",
                        "visual_intent": "process flow",
                        "transition": "fade",
                        "citations": ["S1"],
                    },
                    {
                        "id": "scene-4",
                        "duration_seconds": 3,
                        "kicker": "04",
                        "title": "The takeaway",
                        "on_screen_text": "Reliable research",
                        "narration": (
                            "The final result is useful, inspectable, and source grounded."
                        ),
                        "media_query": "",
                        "visual_intent": "closing statement",
                        "transition": "fade",
                        "citations": ["S1"],
                    },
                ],
            },
        )
    )
    service = VisualArtifactService(
        artifacts=store,
        presentations=PresentationRenderer(),
        videos=VideoOverviewRenderer(),
    )

    result = service.materialize(artifact)

    assert result.status == "ready", result.error
    direction = result.data["direction"]
    assert direction.get("media_policy", direction.get("mediaPolicy")) == "video-first"
    assert result.data["scenes"][0]["visual"]["layout"] == "hero_cinematic"
    for name in ("video_filename", "timeline_filename"):
        assert store.asset_path(
            artifact.workspace_id, artifact.artifact_id, result.data[name]
        ).is_file()
    if include_captions:
        assert store.asset_path(
            artifact.workspace_id,
            artifact.artifact_id,
            result.data["captions_filename"],
        ).is_file()
    else:
        assert "captions_filename" not in result.data
        assert not store.asset_path(
            artifact.workspace_id,
            artifact.artifact_id,
            "captions.srt",
        ).exists()


def test_visual_artifacts_keeps_the_media_asset_compatibility_export() -> None:
    assert MediaAsset is CanonicalMediaAsset
